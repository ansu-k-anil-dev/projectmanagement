from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from datetime import datetime

from auth_module.models import Coridinator, Student
from .models import StudentGroup, Evaluation
from student_module.models import TopicSubmission,PresentationSubmission,RecordSubmission, ProjectSubmission
from .forms import StudentGroupForm, EvaluationForm

@login_required(login_url='login_page')
def view_students(request):
    students = Student.objects.all()
    user_role = Coridinator.objects.get(admin_id=request.user.id).role
    print(user_role)
    exclude_student_ids = StudentGroup.objects.values_list('student__id', flat=True).distinct()
    if exclude_student_ids:
        students = Student.objects.exclude(id__in=exclude_student_ids)
    form = StudentGroupForm()
    context={
        'students':students,
        'form': form,
        'user_role': user_role
    }
    return render(request,'Cordinator/view_students.html', context)


@login_required(login_url='login_page')
def create_student_group(request):
    if request.method == "POST":
        form = StudentGroupForm(request.POST)
        student_ids = [id.strip() for id in request.POST.get('students_ids', '').split(',')]
        
        if not student_ids or student_ids == ['']:
            messages.error(request, 'Please select at least one student')
            return redirect('view_students')
        
        if form.is_valid():
            student_group = form.save(commit=False)
            
            coordinator_admin = request.user.id
            student_group.cordinator = Coridinator.objects.get(admin_id=coordinator_admin)
            student_group.save()
            
            try:
                students = Student.objects.filter(id__in=student_ids)
                if not students.exists():
                    messages.error(request, 'No valid students selected')
                    student_group.delete()
                    return redirect('view_students')
                
                student_group.student.set(students)
                messages.success(request, f'Student group created with {students.count()} students!')
                return redirect('view_student_group')
            
            except Exception as e:
                messages.error(request, f'Error adding students: {str(e)}')
                student_group.delete()  # Clean up on error
                return redirect('view_students')
        else:
            messages.error(request, 'Invalid form data')
        return redirect('view_plan')
    return render(request,'Cordinator/create_student_group.html')



@login_required(login_url='login_page')
def view_student_group(request):
    student_groups = StudentGroup.objects.all()
    user_role = Coridinator.objects.get(admin_id=request.user.id).role
    faculties = Coridinator.objects.filter(role='Faculty')
    evaluation_form = EvaluationForm()
    evaluation_data = Evaluation.objects.filter(cordinator=request.user.id)
    context = {
        'student_groups': student_groups,
        'user_role': user_role,
        'faculties': faculties,
        'evaluation_form': evaluation_form,
        'evaluation_data': evaluation_data
    }
    return render(request,'Cordinator/view_student_group.html', context)


@login_required(login_url='login_page')
def assign_faculty(request):
    if request.method=="POST":
        group_id = request.POST.get('group_id')
        faculty_id = request.POST.get('faculty_id')
        
        student_group_obj = StudentGroup.objects.get(id=group_id)
        student_group_obj.assigned_faculty_id = int(faculty_id)
        student_group_obj.save()
        messages.success(request, f'Faculty has been Assigned to Group {group_id}!')
        return redirect('view_student_group')

@login_required(login_url='login_page')
def assign_obtained_mark(request):
    coordinator_admin = request.user.id
    if request.method == "POST":
        group_id = request.POST.get('group_id')
        student_group_obj = StudentGroup.objects.get(id=group_id)
        project_title = TopicSubmission.objects.get(student_group=student_group_obj).title
        if request.method == 'POST':
            form = EvaluationForm(request.POST)
            if form.is_valid():
                evaluation = form.save(commit=False)
                evaluation.project_title = project_title
                evaluation.cordinator = Coridinator.objects.get(admin_id=coordinator_admin)
                evaluation.student_group = student_group_obj  # Get this from your view context
                evaluation.save()
                student_group_obj.marks_obtained = evaluation.internal_mark
                student_group_obj.save()
            messages.success(request, f'Mark has been Assigned to Group {group_id}!')
        else:
            messages.error(request, 'Invalid form data')
        return redirect('view_student_group')


@login_required(login_url='login_page')
def add_student_to_group(request, group_id):
    group_id = group_id
    if request.method == "POST":
        group_id = request.POST.get('group_id')
        student_ids = [id.strip() for id in request.POST.get('students_ids', '').split(',')]
        
        if not student_ids or student_ids == ['']:
            messages.error(request, 'Please select at least one student')
            return redirect('add_student_to_group')
        else:
            student_group = StudentGroup.objects.get(id=group_id)
            existing_student_ids = student_group.student.values_list('id', flat=True)  # Get existing student IDs from the group
            all_student_ids = list(set(student_ids).union(existing_student_ids))
            students = Student.objects.filter(id__in=all_student_ids)
            student_group.student.set(students)
            messages.success(request, f'Student group updated with {students.count()} students!')
            return redirect('view_student_group')
    else:
        exclude_student_ids = StudentGroup.objects.values_list('student__id', flat=True).distinct()
        student_objs = Student.objects.exclude(id__in=exclude_student_ids)
        context={
            'students':student_objs,
            'group_id': group_id
        }
        return render(request,'Cordinator/add_student_to_group.html', context)


@login_required(login_url='login_page')
def remove_student_from_group(request,group_id,student_id):
    student_group = StudentGroup.objects.get(id=group_id)
    student = Student.objects.get(id=student_id)
    if student in student_group.student.all():
        student_group.student.remove(student)
        messages.success(request, f'Removed {student.admin.first_name} {student.admin.last_name} from group.')
    else:
        messages.error(request, 'Student not found in this group.')
    return redirect('view_student_group')


@login_required(login_url='login_page')
def edit_student_group(request, id):
    student_group = StudentGroup.objects.get(id=id)
    form = StudentGroupForm(instance=student_group)
    context = {
        'form': form,
        'group_id': id
    }
    return render(request, 'Cordinator/edit_student_group.html', context)


@login_required(login_url='login_page')
def update_student_group(request):
    if request.method=="POST":
        group_id = request.POST.get('group_id')
        student_group = StudentGroup.objects.get(id=group_id)
        form = StudentGroupForm(request.POST, instance=student_group)
        if form.is_valid():
            group = form.save()
            messages.success(request, 'Student group updated successfully!')
            return redirect('view_student_group')
        else:
            messages.error(request, 'Invalid form data')
            return redirect('view_student_group')
    else:
        messages.error(request, 'Not Updated')
        return redirect('view_student_group')


@login_required(login_url='login_page')
def delete_student_group(request,id):
    student_group = StudentGroup.objects.filter(id=id)
    student_group.delete()
    messages.success(request, 'Gorup has been successfully Deleted !')
    return redirect('view_student_group')


@login_required(login_url='login_page')
def view_topic_cordinator(request):
    if request.method == "POST":
        topic_id = request.POST.get('topic_id')
        remark = request.POST.get('remark')
        topic_obj = TopicSubmission.objects.get(id=topic_id)
        
        topic_obj.remarks = remark
        topic_obj.status = 'viewed'
        topic_obj.save()

        messages.success(request, 'Topic Remark has been successfully Added !')
        return redirect('view_topic_cordinator')
    else:
        user_role = Coridinator.objects.get(admin_id=request.user.id).role
        if user_role == 'Faculty':
            cordinator = Coridinator.objects.filter(admin=request.user).first()
            student_groups = StudentGroup.objects.filter(assigned_faculty_id=cordinator.admin_id)
        else:
            student_groups = StudentGroup.objects.all()
        topic_data = TopicSubmission.objects.filter(student_group__in = student_groups)
        
        context = {
            'topic_data': topic_data,
            'user_role': user_role
        }
        
        return render(request,'Cordinator/view_topic_cordinator.html', context)
   
 
@login_required(login_url='login_page')
def view_presentation_cordinator(request):
    if request.method == "POST":
        id = request.POST.get('id')
        remark = request.POST.get('remark')
        obj = PresentationSubmission.objects.get(id=id)
        
        obj.remarks = remark
        obj.status = 'viewed'
        obj.save()

        messages.success(request, 'Remark has been successfully Added !')
        return redirect('view_presentation_cordinator')
    else:
        user_role = Coridinator.objects.get(admin_id=request.user.id).role
        if user_role == 'Faculty':
            cordinator = Coridinator.objects.filter(admin=request.user).first()
            student_groups = StudentGroup.objects.filter(assigned_faculty_id=cordinator.admin_id)
        else:
            student_groups = StudentGroup.objects.all()
        data = PresentationSubmission.objects.filter(student_group__in = student_groups)
        
        context = {
            'data': data,
            'user_role': user_role
        }
        
        return render(request,'Cordinator/view_presentation_cordinator.html', context)


@login_required(login_url='login_page')
def view_record_cordinator(request):
    if request.method == "POST":
        id = request.POST.get('id')
        remark = request.POST.get('remark')
        obj = RecordSubmission.objects.get(id=id)
        
        obj.remarks = remark
        obj.status = 'viewed'
        obj.save()

        messages.success(request, 'Remark has been successfully Added !')
        return redirect('view_record_cordinator')
    else:
        user_role = Coridinator.objects.get(admin_id=request.user.id).role
        if user_role == 'Faculty':
            cordinator = Coridinator.objects.filter(admin=request.user).first()
            student_groups = StudentGroup.objects.filter(assigned_faculty_id=cordinator.admin_id)
        else:
            student_groups = StudentGroup.objects.all()
        data = RecordSubmission.objects.filter(student_group__in = student_groups)
        
        context = {
            'data': data
        }
        
        return render(request,'Cordinator/view_record_cordinator.html', context)


@login_required(login_url='login_page')
def view_project_cordinator(request):
    if request.method == "POST":
        id = request.POST.get('id')
        remark = request.POST.get('remark')
        obj = ProjectSubmission.objects.get(id=id)
        
        obj.remarks = remark
        obj.status = 'viewed'
        obj.save()

        messages.success(request, 'Remark has been successfully Added !')
        return redirect('view_project_cordinator')
    else:
        user_role = Coridinator.objects.get(admin_id=request.user.id).role
        if user_role == 'Faculty':
            cordinator = Coridinator.objects.filter(admin=request.user).first()
            student_groups = StudentGroup.objects.filter(assigned_faculty_id=cordinator.admin_id)
        else:
            student_groups = StudentGroup.objects.all()
        data = ProjectSubmission.objects.filter(student_group__in = student_groups)
        
        context = {
            'data': data,
            'user_role': user_role
        }
        
        return render(request,'Cordinator/view_project_cordinator.html', context)


@login_required(login_url='login_page')
def download_evaluations_excel(request):
    # Create a new workbook and select the active sheet
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Evaluations"

    # Define headers
    headers = [
        'Group No.',
        'Students',
        'Assigned Faculty',
        'Project Title',
        'Viva Mark',
        'Demo Mark',
        'Presentation Mark',
        'Internal Mark',
        'Internal Grade',
        'Remark',
        'Added On'
    ]

    # Style for headers
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    
    # Write headers
    for col, header in enumerate(headers, 1):
        cell = sheet.cell(row=1, column=col)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')

    # Fetch all student groups and their evaluations
    student_groups = StudentGroup.objects.all()
    row = 2

    for group in student_groups:
        evaluation = Evaluation.objects.filter(student_group=group).first()
        if evaluation:
            # Get faculty name
            faculty_name = "Not Assigned"
            if group.assigned_faculty_id:
                faculty = Coridinator.objects.filter(admin_id=group.assigned_faculty_id).first()
                if faculty:
                    faculty_name = f"{faculty.admin.first_name} {faculty.admin.last_name}"

            # Get students names
            students_names = ", ".join([
                f"{student.admin.first_name} {student.admin.last_name}"
                for student in group.student.all()
            ])

            # Write data
            data = [
                f"Group {group.id}",
                students_names,
                faculty_name,
                evaluation.project_title,
                evaluation.viva_mark or "N/A",
                evaluation.demo_mark or "N/A",
                evaluation.presentation_mark or "N/A",
                evaluation.internal_mark,
                evaluation.internal_grade,
                evaluation.remark,
                evaluation.added_on.strftime("%d-%m-%Y %H:%M")
            ]

            for col, value in enumerate(data, 1):
                cell = sheet.cell(row=row, column=col)
                cell.value = value
                cell.alignment = Alignment(horizontal='center' if isinstance(value, (int, float)) else 'left')

            row += 1

    # Adjust column widths
    for column in sheet.columns:
        max_length = 0
        column_letter = openpyxl.utils.get_column_letter(column[0].column)
        
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        
        adjusted_width = (max_length + 2)
        sheet.column_dimensions[column_letter].width = adjusted_width

    # Create the HttpResponse with Excel content type
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename=evaluations_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'

    # Save the workbook to the response
    workbook.save(response)
    return response